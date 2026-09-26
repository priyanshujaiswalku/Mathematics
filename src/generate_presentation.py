"""
Script to generate a professional Mid-Term Dissertation Presentation (.pptx)
for Priyanshu Kumar (Department of Mathematics).
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_midterm_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    
    # Color palette
    NAVY = RGBColor(16, 44, 87)       # #102C57
    BLUE = RGBColor(53, 95, 142)      # #355F8E
    DARK_GRAY = RGBColor(50, 50, 50)
    LIGHT_BG = RGBColor(245, 247, 250)
    WHITE = RGBColor(255, 255, 255)
    GOLD = RGBColor(212, 160, 23)
    
    blank_layout = prs.slide_layouts[6]
    
    def add_header(slide, title_text, category="MID-TERM DISSERTATION REVIEW"):
        # Header banner shape
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = NAVY
        header_box.line.color.rgb = NAVY
        
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD
        
        # Title text
        tf = header_box.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.LEFT
        header_box.text_frame.margin_left = Inches(0.8)
        header_box.text_frame.margin_top = Inches(0.35)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY
    
    # Tag
    tb_tag = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.5))
    p = tb_tag.text_frame.paragraphs[0]
    p.text = "FINAL YEAR M.SC. MATHEMATICS DISSERTATION | MID-TERM PROGRESS EVALUATION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GOLD
    
    # Main Title
    tb_title = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.2))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = "Physics-Informed Neural Networks for Forward and Inverse Problems in Nonlinear PDEs"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    
    # Subtitle
    tb_sub = slide1.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.3), Inches(0.8))
    p_sub = tb_sub.text_frame.paragraphs[0]
    p_sub.text = "A Scientific Machine Learning Framework for Mesh-Free Numerical Solutions & Parameter Discovery"
    p_sub.font.size = Pt(18)
    p_sub.font.italic = True
    p_sub.font.color.rgb = RGBColor(200, 215, 235)
    
    # Author & Info Box
    tb_info = slide1.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.5))
    tf_info = tb_info.text_frame
    p1 = tf_info.paragraphs[0]
    p1.text = "Candidate: Priyanshu Kumar"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tf_info.add_paragraph()
    p2.text = "Department of Mathematics | Computational Mathematics & SciML Research"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(180, 200, 220)

    # -------------------------------------------------------------
    # SLIDE 2: Research Motivation & Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Research Motivation & Problem Statement")
    
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf2 = tb2.text_frame
    
    def add_bullet(tf, title, body):
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = "• " + title + ": "
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = NAVY
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)
        
    add_bullet(tf2, "The Core Challenge in Scientific Computing",
               "Partial Differential Equations (PDEs) govern heat transfer, fluid flow, and quantum mechanics. Traditional numerical methods (FDM, FEM, FVM) are fundamentally grid-dependent.")
    add_bullet(tf2, "Limitations of Classical Methods",
               "Mesh generation is computationally expensive, scales poorly in higher dimensions (Curse of Dimensionality), and struggles with complex boundary geometry and sparse/noisy sensor data.")
    add_bullet(tf2, "The SciML Paradigm (Physics-Informed Deep Learning)",
               "Instead of treating neural networks as pure black-box curve fitters, Physics-Informed Neural Networks (PINNs) enforce the governing differential equations directly as inductive bias.")
    add_bullet(tf2, "Mid-Term Objective",
               "Independently re-implement foundational PINN algorithms in modern PyTorch, conduct rigorous comparative benchmarks against classical Finite Difference Method (FDM), and prepare for inverse parameter discovery.")

    # -------------------------------------------------------------
    # SLIDE 3: Mathematical Formulation of PINNs
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Mathematical Framework & Composite Loss Formulation")
    
    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf3 = tb3.text_frame
    
    add_bullet(tf3, "General PDE Operator in Hilbert Space",
               "Let D[u](x, t) = u_t + N_x[u] = 0 in domain Q_T, with boundary trace B[u] = 0 on Sigma_T and initial Cauchy data u(x, 0) = u_0(x).")
    add_bullet(tf3, "Universal Derivative Approximation Theorem (Pinkus, 1999)",
               "Smooth neural networks sigma in C^infinity (such as tanh or sin) can simultaneously approximate both the solution field u and its partial derivatives to arbitrary precision.")
    add_bullet(tf3, "Reverse-Mode Automatic Differentiation (Autograd)",
               "Exact partial derivatives (u_t, u_x, u_xx) are evaluated using the chain rule on the computational graph — eliminating numerical truncation error O(dx^2) without any spatial mesh.")
    add_bullet(tf3, "Variational Loss Functional via Monte Carlo Quadrature",
               "L(theta) = lambda_data * MSE_data + lambda_pde * MSE_pde, where MSE_pde approximates the continuous L^2-norm integral ||D[u_hat]||^2 over scattered interior collocation points.")

    # -------------------------------------------------------------
    # SLIDE 3B: Computational Architecture Diagram
    # -------------------------------------------------------------
    slide3b = prs.slides.add_slide(blank_layout)
    add_header(slide3b, "3. Computational Graph & Physics-Informed Workflow")
    
    diag_path = "reports/figures/pinn_mathematical_architecture.png"
    if os.path.exists(diag_path):
        slide3b.shapes.add_picture(diag_path, Inches(0.8), Inches(1.3), width=Inches(11.73))

    # -------------------------------------------------------------
    # SLIDE 4: Experiment 1: 1D Poisson Benchmark (PINN vs FDM)
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Benchmark 1: 1D Poisson Equation (PINN vs Classical FDM)")
    
    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.5))
    tf4 = tb4.text_frame
    p = tf4.paragraphs[0]
    p.text = "Mathematical Formulation: -d^2u/dx^2 = pi^2 * sin(pi * x),  x in [-1, 1], with u(-1) = u(1) = 0\nAnalytical Solution: u*(x) = sin(pi * x)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE
    
    # Insert Figure
    fig_path = "reports/figures/poisson_pinn_vs_fdm.png"
    if os.path.exists(fig_path):
        slide4.shapes.add_picture(fig_path, Inches(0.8), Inches(2.3), width=Inches(11.73))

    # -------------------------------------------------------------
    # SLIDE 5: Quantitative Comparison Table (FDM vs PINN)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Quantitative Error & Performance Analysis")
    
    # Add Table
    rows, cols = 5, 4
    left = Inches(1.2)
    top = Inches(1.8)
    width = Inches(10.9)
    height = Inches(3.2)
    
    table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    headers = ["Evaluation Metric", "Classical FDM (N=50 Grid)", "PINN (PyTorch Autograd)", "Mathematical Takeaway"]
    data = [
        ["Relative L2 Error", "1.3713 × 10^-3", "7.0039 × 10^-4", "PINN achieved 49% lower error"],
        ["Max Absolute Error", "1.5204 × 10^-3", "7.4602 × 10^-4", "Smooth pointwise residual distribution"],
        ["Mesh Requirement", "Strict equidistant grid", "Mesh-free (scattered points)", "PINN avoids meshing bottlenecks"],
        ["Computation / Training", "Linear system solve (tridiagonal)", "Adam Optimization (3.23s CPU)", "Fast convergence with Tanh activation"]
    ]
    
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.size = Pt(14)
            p.alignment = PP_ALIGN.CENTER
            
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 == 0 else LIGHT_BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.color.rgb = DARK_GRAY
                if c == 2 and r == 0:
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(0, 128, 0)

    # -------------------------------------------------------------
    # SLIDE 6: Experiment 2: 1D Viscous Burgers' Equation (Shock Wave)
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Benchmark 2: 1D Viscous Burgers' Equation (Nonlinear Dynamics)")
    
    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.2))
    tf6 = tb6.text_frame
    p = tf6.paragraphs[0]
    p.text = "Governing Equation: u_t + u * u_x - nu * u_xx = 0,  x in [-1, 1], t in [0, 1],  nu = 0.01 / pi\nInitial Condition: u(x, 0) = -sin(pi * x),  Boundary Condition: u(-1, t) = u(1, t) = 0"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE
    
    # Insert Figure
    fig_burgers = "reports/figures/burgers_pinn_solution.png"
    if os.path.exists(fig_burgers):
        slide6.shapes.add_picture(fig_burgers, Inches(0.8), Inches(2.2), width=Inches(11.73))

    # -------------------------------------------------------------
    # SLIDE 7: Summary of Work Completed (Months 1 & 2)
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Mid-Term Achievements & Deliverables")
    
    tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf7 = tb7.text_frame
    
    add_bullet(tf7, "Rigorous Literature Review",
               "Deep theoretical analysis of Raissi et al. (2017) Part I (Forward Solvers) and Part II (Data-Driven Discovery of PDEs).")
    add_bullet(tf7, "Modern Codebase Re-implementation",
               "Built an independent, modular PyTorch SciML repository from scratch, abandoning legacy TensorFlow 1.x dependencies.")
    add_bullet(tf7, "Two Validated Computational Benchmarks",
               "1D Poisson Boundary Value Problem (validated against analytical & FDM) and 1D Nonlinear Burgers' Equation (capturing shock wave formation).")
    add_bullet(tf7, "Formal Research Repository & Documentation",
               "Structured GitHub repository, sprint schedule (10-day cycles), comprehensive progress reports, and mathematical notes.")

    # -------------------------------------------------------------
    # SLIDE 8: Roadmap for Remaining Months (Months 3–6)
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. Research Roadmap for Remaining Months (3 to 6)")
    
    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf8 = tb8.text_frame
    
    add_bullet(tf8, "Month 3: Inverse Problem & Parameter Discovery",
               "Treat physical viscosity nu as a learnable parameter. Test discovery accuracy from sparse data corrupted with 1%, 5%, and 10% Gaussian noise.")
    add_bullet(tf8, "Month 4: Coupled & Complex Systems (Schrödinger Equation)",
               "Extend PINNs to complex-valued wave functions (soliton wave dynamics) and evaluate discrete-time Runge-Kutta PINN schemes.")
    add_bullet(tf8, "Month 5: Optimization & Convergence Enhancement",
               "Implement hybrid Adam -> L-BFGS optimization routine and evaluate self-adaptive loss weighting schemes to resolve gradient pathologies.")
    add_bullet(tf8, "Month 6: Thesis Drafting & Final Defense Preparation",
               "Complete LaTeX dissertation report, generate publication-grade figures, and prepare the final viva presentation.")

    # -------------------------------------------------------------
    # SLIDE 9: Conclusion & Q&A
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    bg9 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg9.fill.solid()
    bg9.fill.fore_color.rgb = NAVY
    bg9.line.color.rgb = NAVY
    
    tb9 = slide9.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.0))
    tf9 = tb9.text_frame
    p = tf9.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf9.add_paragraph()
    p2.text = "Questions & Suggestions are Warmly Welcomed."
    p2.font.size = Pt(22)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)
    
    p3 = tf9.add_paragraph()
    p3.text = "Candidate: Priyanshu Kumar | Department of Mathematics"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(180, 200, 220)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(25)

    os.makedirs("presentations/midterm_presentation", exist_ok=True)
    out_path = "presentations/midterm_presentation/Midterm_PINNs_Presentation.pptx"
    prs.save(out_path)
    print(f"[SUCCESS] Presentation saved to: {out_path}")

if __name__ == "__main__":
    create_midterm_presentation()
